#!/usr/bin/env python3
import argparse
import csv
import re
from copy import copy
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class Product:
    product_id: str
    name: str
    amount_fen: str

    @property
    def cn_price(self) -> str:
        return f"{Decimal(self.amount_fen) / Decimal(100):.2f}"


def parse_product_line(line: str) -> Product:
    line = line.strip()
    if not line:
        raise ValueError("Empty product line")
    parts = [part for part in re.split(r"\t+|\s{2,}", line) if part]
    if len(parts) < 3:
        parts = line.split()
    if len(parts) < 3:
        raise ValueError(f"Cannot parse product line: {line}")
    product_id = parts[0].strip()
    cash_token = parts[-1].strip()
    name = " ".join(parts[1:-1]).strip()
    match = re.search(r"Cash\s*:\s*(\d+)", cash_token, re.IGNORECASE)
    if not match:
        raise ValueError(f"Missing Cash price in line: {line}")
    if not re.match(r"^[A-Za-z0-9][A-Za-z0-9_.]*$", product_id):
        raise ValueError(f"Invalid product ID for Huawei: {product_id}")
    if "|" in name:
        raise ValueError(f"Product name cannot contain '|': {name}")
    return Product(product_id=product_id, name=name, amount_fen=match.group(1))


def load_products(args) -> list[Product]:
    lines = []
    if args.products:
        lines.extend(args.products)
    if args.products_file:
        lines.extend(Path(args.products_file).read_text(encoding="utf-8-sig").splitlines())
    products = [parse_product_line(line) for line in lines if line.strip()]
    if not products:
        raise ValueError("No products provided")
    seen = set()
    duplicates = []
    for product in products:
        if product.product_id in seen:
            duplicates.append(product.product_id)
        seen.add(product.product_id)
    if duplicates:
        raise ValueError(f"Duplicate product IDs in request: {duplicates}")
    return products


def cn_price_from_huawei(price_cell) -> str | None:
    parts = str(price_cell or "").split(";")
    for index, part in enumerate(parts[:-1]):
        if part == "CN":
            return parts[index + 1]
    return None


def put_cn_price_first(price_cell) -> str:
    parts = str(price_cell or "").split(";")
    if len(parts) % 2:
        raise ValueError("Huawei Price has an odd number of country/price tokens")
    pairs = list(zip(parts[0::2], parts[1::2]))
    cn_pairs = [pair for pair in pairs if pair[0] == "CN"]
    if len(cn_pairs) != 1:
        raise ValueError(f"Huawei Price must contain exactly one CN entry, found {len(cn_pairs)}")
    return ";".join(value for pair in cn_pairs + [pair for pair in pairs if pair[0] != "CN"] for value in pair)


def copy_cell_style(source, target):
    target._style = copy(source._style)
    target.number_format = source.number_format
    target.alignment = copy(source.alignment)
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)


def make_huawei(products: list[Product], huawei_export: Path, output: Path, price_mode: str = "full"):
    wb = load_workbook(huawei_export)
    ws = wb.active
    existing_ids = set()
    refs: dict[str, tuple[list, list]] = {}

    for row_idx in range(3, ws.max_row + 1):
        product_id = str(ws.cell(row_idx, 1).value or "")
        if product_id:
            existing_ids.add(product_id)
        price = cn_price_from_huawei(ws.cell(row_idx, 4).value)
        if price and price not in refs:
            values = [ws.cell(row_idx, col).value for col in range(1, ws.max_column + 1)]
            styles = [copy(ws.cell(row_idx, col)._style) for col in range(1, ws.max_column + 1)]
            refs[price] = (values, styles)

    duplicate_ids = [product.product_id for product in products if product.product_id in existing_ids]
    if duplicate_ids:
        raise ValueError(f"Huawei product IDs already exist: {duplicate_ids}")

    missing_prices = sorted({product.cn_price for product in products if product.cn_price not in refs})
    if missing_prices:
        raise ValueError(f"No Huawei same-price reference rows for CN prices: {missing_prices}")

    if ws.max_row > 2:
        ws.delete_rows(3, ws.max_row - 2)

    for out_row, product in enumerate(products, start=3):
        values, styles = refs[product.cn_price]
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(out_row, col)
            cell._style = copy(styles[col - 1])
            cell.value = values[col - 1]
        ws.cell(out_row, 1).value = product.product_id
        ws.cell(out_row, 2).value = "0"
        ws.cell(out_row, 3).value = f"zh_CN|{product.name}|{product.name}"
        if price_mode == "cn-only":
            ws.cell(out_row, 4).value = f"CN;{product.cn_price}"
        else:
            ws.cell(out_row, 4).value = put_cn_price_first(values[3])
        ws.cell(out_row, 5).value = ""
        ws.cell(out_row, 6).value = ""
        ws.cell(out_row, 7).value = "0"
        ws.cell(out_row, 9).value = "1"
        if ws.max_column >= 10:
            ws.cell(out_row, 10).value = ""

    wb.save(output)


def read_csv(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return reader.fieldnames, list(reader)


def make_gm_csv(products: list[Product], source: Path, output: Path, backend_name: str):
    fieldnames, rows = read_csv(source)
    existing_ids = {row["物品ID"] for row in rows}
    duplicate_ids = [product.product_id for product in products if product.product_id in existing_ids]
    if duplicate_ids:
        raise ValueError(f"{backend_name} product IDs already exist: {duplicate_ids}")

    refs = {}
    for row in rows:
        amount = row.get("物品金额（分）", "")
        if amount:
            refs[amount] = row

    missing_amounts = sorted({product.amount_fen for product in products if product.amount_fen not in refs})
    if missing_amounts:
        raise ValueError(f"No {backend_name} same-price reference rows for amounts: {missing_amounts}")

    output_rows = []
    for product in products:
        row = dict(refs[product.amount_fen])
        row["物品ID"] = product.product_id
        row["物品名称"] = product.name
        row["物品金额（分）"] = product.amount_fen
        if "创建时间" in row:
            row["创建时间"] = ""
        output_rows.append(row)

    with output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, lineterminator="\n")
        writer.writeheader()
        writer.writerows(output_rows)


def verify_huawei(path: Path, products: list[Product], price_mode: str = "full"):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    rows = [row for row in rows if row and row[0]]
    if len(rows) != len(products):
        raise ValueError(f"Huawei row count mismatch: {len(rows)} != {len(products)}")
    by_id = {str(row[0]): row for row in rows}
    for product in products:
        row = by_id.get(product.product_id)
        if not row:
            raise ValueError(f"Missing Huawei row: {product.product_id}")
        if row[2] != f"zh_CN|{product.name}|{product.name}":
            raise ValueError(f"Huawei name mismatch: {product.product_id}")
        price_parts = str(row[3]).split(";")
        if len(price_parts) % 2:
            raise ValueError(f"Huawei price token count is odd: {product.product_id}")
        if price_parts[0] != "CN" or price_parts[1] != product.cn_price:
            raise ValueError(f"Huawei base price is not CN: {product.product_id}")
        if price_mode == "cn-only" and len(price_parts) != 2:
            raise ValueError(f"Huawei CN-only price contains extra countries: {product.product_id}")
        price_map = dict(zip(price_parts[0::2], price_parts[1::2]))
        if price_map.get("CN") != product.cn_price:
            raise ValueError(f"Huawei CN price mismatch: {product.product_id}")


def verify_csv(path: Path, products: list[Product], backend_name: str):
    _, rows = read_csv(path)
    if len(rows) != len(products):
        raise ValueError(f"{backend_name} row count mismatch: {len(rows)} != {len(products)}")
    by_id = {row["物品ID"]: row for row in rows}
    for product in products:
        row = by_id.get(product.product_id)
        if not row:
            raise ValueError(f"Missing {backend_name} row: {product.product_id}")
        if row["物品名称"] != product.name or row["物品金额（分）"] != product.amount_fen:
            raise ValueError(f"{backend_name} field mismatch: {product.product_id}")


def compact_name(products: list[Product]) -> str:
    if len(products) == 1:
        return products[0].product_id
    return f"{products[0].product_id}_{products[-1].product_id}_{len(products)}items"


def main():
    parser = argparse.ArgumentParser(description="Generate Huawei, CN, and HMT billing batch files.")
    parser.add_argument("--products", action="append", help="Product line, e.g. '777026\\t镜中迷踪\\tCash:1800'. Can be repeated.")
    parser.add_argument("--products-file", help="UTF-8 text file with one product line per row.")
    parser.add_argument("--huawei-export", required=True)
    parser.add_argument("--cn-export", required=True)
    parser.add_argument("--hmt-export", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument(
        "--huawei-price-mode",
        choices=("full", "cn-only"),
        default="full",
        help="Use 'full' to preserve all localized prices or 'cn-only' to let Huawei fill missing countries when creating products.",
    )
    args = parser.parse_args()

    products = load_products(args)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    suffix = compact_name(products)

    huawei_output = output_dir / f"华为_批量新增_{suffix}.xlsx"
    cn_output = output_dir / f"国服_批量新增_{suffix}.csv"
    hmt_output = output_dir / f"港澳台_批量新增_{suffix}.csv"

    make_huawei(products, Path(args.huawei_export), huawei_output, args.huawei_price_mode)
    make_gm_csv(products, Path(args.cn_export), cn_output, "CN GM")
    make_gm_csv(products, Path(args.hmt_export), hmt_output, "HMT GM")

    verify_huawei(huawei_output, products, args.huawei_price_mode)
    verify_csv(cn_output, products, "CN GM")
    verify_csv(hmt_output, products, "HMT GM")

    print(huawei_output)
    print(cn_output)
    print(hmt_output)


if __name__ == "__main__":
    main()
